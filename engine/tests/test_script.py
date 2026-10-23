"""Etapa de script Python do usuário (API `import datahub`)."""

from __future__ import annotations

from conftest import criar_xlsx, ler_xlsx

SCRIPT_QUE_FILTRA = """
import datahub
from openpyxl import Workbook, load_workbook

ctx = datahub.contexto()
print("começando")
linhas = list(load_workbook(ctx["entrada"]).active.iter_rows(values_only=True))
datahub.progresso(50, "metade")
saida = Workbook()
for linha in [linhas[0]] + [l for l in linhas[1:] if l[0] > 1]:
    saida.active.append(linha)
saida.save(ctx["saida"])
datahub.log("filtro aplicado", "aviso")
"""


def _etapa_script(caminho, pasta=None):
    return {"ferramenta": "script-python", "script": {"principal": str(caminho), "pasta": pasta}}


def test_script_le_a_etapa_grava_a_saida_e_registra_logs(tmp_path, executar):
    script = tmp_path / "filtro.py"
    script.write_text(SCRIPT_QUE_FILTRA, encoding="utf-8")
    base = criar_xlsx(tmp_path / "b.xlsx", ["N"], [[1], [2], [3]])
    resultado = executar([base], [_etapa_script(script)])

    assert ler_xlsx(resultado.saida)["Dados"] == [["N"], [2], [3]]
    logs = [(e["nivel"], e["mensagem"]) for e in resultado.do_tipo("log")]
    assert logs == [("info", "começando"), ("aviso", "filtro aplicado")]
    assert any(e["mensagem"] == "metade" for e in resultado.do_tipo("progresso"))
    assert resultado.do_tipo("etapa.fim")[0]["linhas_depois"] == 2


def test_script_importa_modulos_da_propria_pasta(tmp_path, executar):
    pasta = tmp_path / "scripts"
    (pasta / "util").mkdir(parents=True)
    (pasta / "util" / "__init__.py").write_text("", encoding="utf-8")
    (pasta / "util" / "texto.py").write_text("MENSAGEM = 'olá do módulo'\n", encoding="utf-8")
    (pasta / "main.py").write_text("from util.texto import MENSAGEM\nprint(MENSAGEM)\n", encoding="utf-8")
    base = criar_xlsx(tmp_path / "b.xlsx", ["N"], [[1]])
    resultado = executar([base], [_etapa_script(pasta / "main.py", str(pasta))])
    assert resultado.do_tipo("log")[0]["mensagem"] == "olá do módulo"


def test_script_que_falha_vira_erro_com_detalhe(tmp_path, executar):
    script = tmp_path / "quebra.py"
    script.write_text("raise ValueError('coluna X ausente')\n", encoding="utf-8")
    base = criar_xlsx(tmp_path / "b.xlsx", ["N"], [[1]])
    resultado = executar([base], [_etapa_script(script)])
    erro = resultado.fim["erro"]
    assert resultado.status == "erro"
    assert erro["mensagem"] == "O script quebra.py falhou: coluna X ausente"
    assert erro["etapa"] == 0 and "ValueError" in erro["detalhe"]


def test_script_sem_saida_mantem_a_planilha_e_avisa(tmp_path, executar):
    script = tmp_path / "nada.py"
    script.write_text("x = 1\n", encoding="utf-8")
    base = criar_xlsx(tmp_path / "b.xlsx", ["N"], [[1]])
    resultado = executar([base], [_etapa_script(script)])
    assert resultado.status == "concluido_com_avisos"
    assert ler_xlsx(resultado.saida)["Dados"] == [["N"], [1]]


def test_script_inexistente(tmp_path, executar):
    base = criar_xlsx(tmp_path / "b.xlsx", ["N"], [[1]])
    resultado = executar([base], [_etapa_script(tmp_path / "sumiu.py")])
    assert resultado.fim["erro"]["mensagem"].startswith("Script Python não encontrado")
