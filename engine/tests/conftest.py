"""Auxiliares dos testes. Todos os dados são fictícios."""

from __future__ import annotations

import io
import json
import random
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from datahub_engine.eventos import Emissor, validar_evento
from datahub_engine.execucao import Cancelamento, executar_plano
from datahub_engine.plano import carregar_plano

CONTRATO = Path(__file__).resolve().parents[1] / "contrato"


def criar_xlsx(caminho: Path, cabecalho: list[str], linhas: list[list[Any]], abas_extras: int = 0) -> Path:
    livro = Workbook()
    folha = livro.active
    folha.title = "Dados"
    folha.append(cabecalho)
    for linha in linhas:
        folha.append(linha)
    for numero in range(abas_extras):
        livro.create_sheet(f"Extra {numero + 1}").append(["x"])
    livro.save(caminho)
    return caminho


def ler_xlsx(caminho: Path) -> dict[str, list[list[Any]]]:
    livro = load_workbook(caminho)
    return {folha.title: [list(linha) for linha in folha.iter_rows(values_only=True)] for folha in livro.worksheets}


class Execucao:
    def __init__(self, status: str, eventos: list[dict[str, Any]], saida: Path) -> None:
        self.status = status
        self.eventos = eventos
        self.saida = saida

    def do_tipo(self, tipo: str) -> list[dict[str, Any]]:
        return [evento for evento in self.eventos if evento["tipo"] == tipo]

    @property
    def fim(self) -> dict[str, Any]:
        return self.do_tipo("execucao.fim")[-1]

    def resumo(self, indice: int = 0) -> dict[str, int]:
        return next(e for e in self.do_tipo("etapa.fim") if e["indice"] == indice)["resumo"]


@pytest.fixture
def executar(tmp_path):
    """Executa um plano em memória e devolve status, eventos (validados) e caminho da saída."""

    def _executar(
        entradas: list[Path],
        etapas: list[dict[str, Any]],
        saida: str = "saida.xlsx",
        cancelar: bool = False,
        semente: int = 7,
    ) -> Execucao:
        destino = tmp_path / saida
        plano = carregar_plano(
            {
                "versao": 1,
                "nome": "Teste",
                "origem": "processo",
                "entradas": [{"caminho": str(caminho)} for caminho in entradas],
                "saida": {"caminho": str(destino)},
                "etapas": etapas,
            }
        )
        fluxo = io.StringIO()
        cancelamento = Cancelamento()
        if cancelar:
            cancelamento.pedir()
        status = executar_plano(plano, Emissor("teste", None, saida=fluxo), cancelamento, random.Random(semente))
        eventos = [json.loads(linha) for linha in fluxo.getvalue().splitlines()]
        for evento in eventos:
            validar_evento(evento)
        assert [evento["seq"] for evento in eventos] == list(range(1, len(eventos) + 1))
        return Execucao(status, eventos, destino)

    return _executar
