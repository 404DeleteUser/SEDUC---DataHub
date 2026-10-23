"""As 5 ferramentas no motor, com as regras combinadas no app. Dados fictícios."""

from __future__ import annotations

import json
from datetime import datetime

from openpyxl import load_workbook

from conftest import criar_xlsx, ler_xlsx


def etapa(ferramenta, **parametros):
    return {"ferramenta": ferramenta, "parametros": {"ferramentaId": ferramenta, **parametros}}


class TestEliminarDuplicatas:
    def test_duas_colunas_iguais_e_duplicata(self, tmp_path, executar):
        base = criar_xlsx(tmp_path / "b.xlsx", ["Sequência", "Número"], [["1234", 4], ["1234", 4], ["1234", 3]])
        resultado = executar([base], [etapa("eliminar-duplicatas", colunasChave=["Sequência", "Número"])])
        assert resultado.status == "concluido"
        assert resultado.resumo() == {"duplicadas_removidas": 1, "grupos_com_duplicatas": 1}
        assert len(ler_xlsx(resultado.saida)["Dados"]) == 3  # cabeçalho + 2 linhas

    def test_espaco_no_fim_caixa_e_acento_diferenciam(self, tmp_path, executar):
        base = criar_xlsx(tmp_path / "b.xlsx", ["Nome"], [["Ana"], ["Ana "], ["ana"], ["Âna"]])
        resultado = executar([base], [etapa("eliminar-duplicatas", colunasChave=["Nome"])])
        assert resultado.resumo()["duplicadas_removidas"] == 0

    def test_linha_que_fica_e_sorteada_mas_reproduzivel_e_na_ordem(self, tmp_path, executar):
        linhas = [["a", 1], ["b", 2], ["a", 3], ["a", 4]]
        base = criar_xlsx(tmp_path / "b.xlsx", ["Chave", "Id"], linhas)
        primeira = executar([base], [etapa("eliminar-duplicatas", colunasChave=["Chave"])], semente=1)
        segunda = executar([base], [etapa("eliminar-duplicatas", colunasChave=["Chave"])], saida="s2.xlsx", semente=1)
        dados = ler_xlsx(primeira.saida)["Dados"][1:]
        assert dados == ler_xlsx(segunda.saida)["Dados"][1:]
        assert len(dados) == 2 and ["b", 2] in dados

    def test_coluna_inexistente_vira_erro_explicado(self, tmp_path, executar):
        base = criar_xlsx(tmp_path / "b.xlsx", ["CPF"], [["1"]])
        resultado = executar([base], [etapa("eliminar-duplicatas", colunasChave=["Matrícula"])])
        assert resultado.status == "erro"
        assert resultado.fim["erro"] == {"mensagem": "A coluna Matrícula não existe na aba Dados.", "etapa": 0}
        assert not resultado.saida.exists()


class TestDicionario:
    def test_contem_sem_caixa_e_acento(self, tmp_path, executar):
        base = criar_xlsx(tmp_path / "b.xlsx", ["Cargo"], [["Professor de Matemática"], ["Diretora"], ["Técnico"]])
        resultado = executar([base], [etapa("dicionario", coluna="Cargo", modo="contem", valores=["PROFESSOR", "tecnico"])])
        assert resultado.resumo() == {"linhas_mantidas": 2, "linhas_removidas": 1}

    def test_exato_e_aviso_quando_nada_corresponde(self, tmp_path, executar):
        base = criar_xlsx(tmp_path / "b.xlsx", ["Cargo"], [["Professor de Matemática"]])
        resultado = executar([base], [etapa("dicionario", coluna="Cargo", modo="exato", valores=["professor"])])
        assert resultado.status == "concluido_com_avisos"
        assert resultado.do_tipo("log")[0]["mensagem"] == "Nenhuma linha corresponde aos valores do dicionário."


class TestMesclar:
    def test_vertical_alinha_colunas_pela_ordem_do_primeiro(self, tmp_path, executar):
        a = criar_xlsx(tmp_path / "a.xlsx", ["CPF", "Nome"], [["1", "Ana"]])
        b = criar_xlsx(tmp_path / "b.xlsx", ["nome", "cpf"], [["Bia", "2"]])
        resultado = executar([a, b], [etapa("mesclar-planilhas", modo="vertical")])
        assert ler_xlsx(resultado.saida)["Dados"] == [["CPF", "Nome"], ["1", "Ana"], ["2", "Bia"]]

    def test_vertical_com_cabecalho_diferente_e_erro(self, tmp_path, executar):
        a = criar_xlsx(tmp_path / "a.xlsx", ["CPF", "Nome"], [])
        b = criar_xlsx(tmp_path / "b.xlsx", ["CPF", "Cargo"], [])
        resultado = executar([a, b], [etapa("mesclar-planilhas", modo="vertical")])
        assert resultado.status == "erro"
        assert "b.xlsx: faltam Nome; sobram Cargo" in resultado.fim["erro"]["mensagem"]

    def test_cada_arquivo_numa_aba(self, tmp_path, executar):
        a = criar_xlsx(tmp_path / "escola.xlsx", ["X"], [[1]])
        b = criar_xlsx(tmp_path / "Escola.xlsx".lower().replace("escola", "turma"), ["Y"], [[2]])
        resultado = executar([a, b], [etapa("mesclar-planilhas", modo="abas")])
        assert list(ler_xlsx(resultado.saida)) == ["escola", "turma"]

    def test_precisa_de_dois_arquivos(self, tmp_path, executar):
        a = criar_xlsx(tmp_path / "a.xlsx", ["X"], [])
        resultado = executar([a], [etapa("mesclar-planilhas", modo="abas")])
        assert resultado.fim["erro"]["mensagem"] == "Mesclar Planilhas precisa de pelo menos dois arquivos de entrada."


class TestDividirPor:
    def test_consolidada_uma_aba_por_valor_e_outros(self, tmp_path, executar):
        base = criar_xlsx(
            tmp_path / "b.xlsx",
            ["Escola", "Aluno"],
            [["Escola A", 1], ["escola a ", 2], ["Escola B", 3], [None, 4], ["Consolidada", 5]],
        )
        resultado = executar([base], [etapa("dividir-por", coluna="Escola")])
        abas = ler_xlsx(resultado.saida)
        assert list(abas) == ["Consolidada", "Escola A", "Escola B", "Consolidada (2)", "Outros"]
        assert len(abas["Consolidada"]) == 6
        assert len(abas["Escola A"]) == 3
        assert resultado.resumo() == {"abas_criadas": 4, "linhas_sem_valor": 1}


class TestZebrar:
    def test_pinta_cabecalho_e_alterna_linhas(self, tmp_path, executar):
        base = criar_xlsx(tmp_path / "b.xlsx", ["X"], [[1], [2], [3]])
        resultado = executar(
            [base], [etapa("zebrar-planilha", corCabecalho="#1f4e79", corLinhaA="#ffffff", corLinhaB="#e8effc")]
        )
        folha = load_workbook(resultado.saida).active
        cores = [folha.cell(linha, 1).fill.fgColor.rgb for linha in range(1, 5)]
        assert cores == ["001F4E79", "00FFFFFF", "00E8EFFC", "00FFFFFF"]
        assert folha.cell(1, 1).font.bold

    def test_cor_invalida_e_erro(self, tmp_path, executar):
        base = criar_xlsx(tmp_path / "b.xlsx", ["X"], [[1]])
        resultado = executar([base], [etapa("zebrar-planilha", corCabecalho="azul", corLinhaA="#ffffff", corLinhaB="#ffffff")])
        assert resultado.fim["erro"]["mensagem"] == "A cor do cabeçalho é inválida (use o formato #RRGGBB)."

    def test_saida_csv_avisa_que_cores_nao_vao(self, tmp_path, executar):
        base = criar_xlsx(tmp_path / "b.xlsx", ["X"], [[1]])
        resultado = executar(
            [base],
            [etapa("zebrar-planilha", corCabecalho="#000000", corLinhaA="#ffffff", corLinhaB="#eeeeee")],
            saida="saida.csv",
        )
        assert resultado.status == "concluido_com_avisos"
        assert resultado.saida.read_text(encoding="utf-8-sig").splitlines() == ["X", "1"]


class TestEntradaESaida:
    def test_numeros_e_datas_continuam_com_o_tipo_original(self, tmp_path, executar):
        data = datetime(2026, 3, 1)
        base = criar_xlsx(tmp_path / "b.xlsx", ["Valor", "Data"], [[1520.5, data]])
        resultado = executar([base], [])
        assert ler_xlsx(resultado.saida)["Dados"][1] == [1520.5, data]

    def test_csv_do_excel_no_windows_com_ponto_e_virgula(self, tmp_path, executar):
        csv = tmp_path / "w.csv"
        csv.write_bytes("Lotação;Nome\nNÚCLEO;Ana\n".encode("cp1252"))
        resultado = executar([csv], [])
        assert resultado.do_tipo("entrada.lida")[0]["colunas"] == 2
        assert ler_xlsx(resultado.saida)["w"] == [["Lotação", "Nome"], ["NÚCLEO", "Ana"]]

    def test_avisa_quando_ignora_abas_extras(self, tmp_path, executar):
        base = criar_xlsx(tmp_path / "b.xlsx", ["X"], [[1]], abas_extras=2)
        resultado = executar([base], [])
        assert resultado.do_tipo("log")[0]["mensagem"] == "Só a primeira aba de b.xlsx foi lida (o arquivo tem 3)."

    def test_arquivo_inexistente_e_erro(self, tmp_path, executar):
        resultado = executar([tmp_path / "nao-existe.xlsx"], [])
        assert resultado.fim["erro"]["mensagem"].startswith("Arquivo não encontrado")

    def test_cancelado_nao_grava_saida(self, tmp_path, executar):
        base = criar_xlsx(tmp_path / "b.xlsx", ["X"], [[1]])
        resultado = executar([base], [], cancelar=True)
        assert resultado.status == "cancelado"
        assert not resultado.saida.exists()


def test_privacidade_nenhum_valor_de_celula_vai_para_os_eventos(tmp_path, executar):
    segredo = "SEGREDO-123"
    linhas = [[f"{segredo}-{i % 3}", f"{segredo}-nome", i] for i in range(9)]
    base = criar_xlsx(tmp_path / "b.xlsx", ["Escola", "Nome", "N"], linhas)
    outra = criar_xlsx(tmp_path / "c.xlsx", ["Escola", "Nome", "N"], linhas)
    resultado = executar(
        [base, outra],
        [
            etapa("mesclar-planilhas", modo="vertical"),
            etapa("eliminar-duplicatas", colunasChave=["Escola", "Nome"]),
            etapa("dicionario", coluna="N", modo="contem", valores=["1", "2"]),
            etapa("dividir-por", coluna="Escola"),
            etapa("zebrar-planilha", corCabecalho="#000000", corLinhaA="#ffffff", corLinhaB="#eeeeee"),
        ],
    )
    assert resultado.status in {"concluido", "concluido_com_avisos"}
    assert segredo not in json.dumps(resultado.eventos, ensure_ascii=False)
