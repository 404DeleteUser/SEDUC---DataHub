"""Pasta de trabalho em memória: abas com um DataFrame cada, lidas e gravadas em disco.

Os DataFrames usam dtype `object`: os valores das células ficam como vieram (número,
data, texto, None), para que a saída preserve os tipos da planilha original. Colunas
com nome repetido são permitidas; as ferramentas acessam colunas por posição.
"""

from __future__ import annotations

import csv
import io
import os
import tempfile
from dataclasses import dataclass, field, replace
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

from .texto import nome_de_aba

EXTENSOES_DE_ENTRADA = {".xlsx", ".xlsm", ".xls", ".csv"}
EXTENSOES_DE_SAIDA = {".xlsx", ".csv"}


class ErroDePlanilha(Exception):
    """Problema com um arquivo, explicado para o usuário."""


def novo_dataframe(cabecalho: list[str], linhas: list[list[Any]]) -> pd.DataFrame:
    return pd.DataFrame(linhas, columns=list(cabecalho), dtype=object)


def valor_python(valor: Any) -> Any:
    """Célula pronta para gravar: NaN/NaT viram None e escalares numpy viram tipos Python."""
    if valor is None:
        return None
    if isinstance(valor, np.generic):
        valor = valor.item()
    if isinstance(valor, float) and np.isnan(valor):
        return None
    if valor is pd.NaT:
        return None
    return valor


@dataclass
class Aba:
    nome: str
    dados: pd.DataFrame

    @classmethod
    def de_linhas(cls, nome: str, cabecalho: list[str], linhas: list[list[Any]]) -> Aba:
        return cls(nome, novo_dataframe(cabecalho, linhas))

    @property
    def cabecalho(self) -> list[str]:
        return [str(coluna) for coluna in self.dados.columns]

    @property
    def linhas(self) -> list[list[Any]]:
        """Cópia das linhas como listas Python (para gravar, comparar e testar)."""
        return [[valor_python(valor) for valor in linha] for linha in self.dados.itertuples(index=False, name=None)]

    @linhas.setter
    def linhas(self, linhas: list[list[Any]]) -> None:
        # Ponte provisória: remover quando todas as ferramentas trabalharem com `dados`.
        self.dados = novo_dataframe(self.cabecalho, linhas)

    def com_dados(self, dados: pd.DataFrame) -> Aba:
        return replace(self, dados=dados)

    def indice_da_coluna(self, coluna: str) -> int | None:
        try:
            return self.cabecalho.index(coluna)
        except ValueError:
            return None


@dataclass
class Zebra:
    cabecalho: str
    linha_a: str
    linha_b: str


@dataclass
class Pasta:
    abas: list[Aba]
    zebra: Zebra | None = None

    @property
    def linhas(self) -> int:
        return sum(len(aba.dados) for aba in self.abas)


@dataclass
class ArquivoLido:
    caminho: Path
    aba: Aba
    total_de_abas: int = 1
    avisos: list[str] = field(default_factory=list)


def _vazia(linha: list[Any]) -> bool:
    return all(valor is None or (isinstance(valor, str) and not valor.strip()) for valor in linha)


def _montar_aba(nome: str, linhas: list[list[Any]], arquivo: str) -> Aba:
    linhas = [list(linha) for linha in linhas]
    while linhas and _vazia(linhas[-1]):
        linhas.pop()
    if not linhas or _vazia(linhas[0]):
        raise ErroDePlanilha(f"Planilha sem cabeçalho: {arquivo}")
    cabecalho = []
    for indice, valor in enumerate(linhas[0]):
        nome_coluna = "" if valor is None else str(valor).strip()
        cabecalho.append(nome_coluna or f"Coluna {indice + 1}")
    largura = len(cabecalho)
    dados = [(linha + [None] * largura)[:largura] for linha in linhas[1:] if not _vazia(linha)]
    return Aba.de_linhas(nome_de_aba(nome), cabecalho, dados)


def _texto_do_csv(bruto: bytes) -> str:
    try:
        return bruto.decode("utf-8-sig")
    except UnicodeDecodeError:
        return bruto.decode("cp1252")


def ler_planilha(caminho: Path) -> ArquivoLido:
    """Lê a primeira aba (como o app faz na prévia)."""
    extensao = caminho.suffix.lower()
    if extensao not in EXTENSOES_DE_ENTRADA:
        raise ErroDePlanilha(f"Formato não suportado: {caminho.name}")
    if not caminho.is_file():
        raise ErroDePlanilha(f"Arquivo não encontrado: {caminho}")

    if extensao == ".csv":
        texto = _texto_do_csv(caminho.read_bytes())
        try:
            dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=";,\t")
        except csv.Error:
            dialeto = csv.excel
        linhas = list(csv.reader(io.StringIO(texto), dialeto))
        return ArquivoLido(caminho, _montar_aba(caminho.stem, linhas, caminho.name))

    if extensao == ".xls":
        import xlrd

        livro = xlrd.open_workbook(str(caminho))
        folha = livro.sheet_by_index(0)
        linhas = [folha.row_values(indice) for indice in range(folha.nrows)]
        return ArquivoLido(caminho, _montar_aba(folha.name, linhas, caminho.name), livro.nsheets)

    from openpyxl import load_workbook

    livro = load_workbook(caminho, read_only=True, data_only=True)
    try:
        folha = livro.worksheets[0]
        linhas = [list(linha) for linha in folha.iter_rows(values_only=True)]
        return ArquivoLido(caminho, _montar_aba(folha.title, linhas, caminho.name), len(livro.worksheets))
    finally:
        livro.close()


def _cor(hexadecimal: str) -> str:
    return hexadecimal.lstrip("#").upper()


def _gravar_xlsx(pasta: Pasta, destino: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    livro = Workbook()
    livro.remove(livro.active)
    for aba in pasta.abas:
        folha = livro.create_sheet(aba.nome)
        folha.append(aba.cabecalho)
        for linha in aba.linhas:
            folha.append(linha)
        if pasta.zebra:
            preenchimentos = [
                PatternFill("solid", fgColor=_cor(pasta.zebra.linha_a)),
                PatternFill("solid", fgColor=_cor(pasta.zebra.linha_b)),
            ]
            for celula in folha[1]:
                celula.fill = PatternFill("solid", fgColor=_cor(pasta.zebra.cabecalho))
                celula.font = Font(bold=True)
            for numero, linha in enumerate(folha.iter_rows(min_row=2), start=0):
                for celula in linha:
                    celula.fill = preenchimentos[numero % 2]
    livro.save(destino)


def _gravar_csv(pasta: Pasta, destino: Path) -> None:
    if len(pasta.abas) != 1:
        raise ErroDePlanilha(
            f"CSV guarda uma aba só, e o resultado tem {len(pasta.abas)}. Escolha a saída em XLSX."
        )
    aba = pasta.abas[0]
    with destino.open("w", encoding="utf-8-sig", newline="") as arquivo:
        escritor = csv.writer(arquivo, delimiter=";")
        escritor.writerow(aba.cabecalho)
        escritor.writerows(["" if valor is None else valor for valor in linha] for linha in aba.linhas)


def gravar_pasta(pasta: Pasta, destino: Path) -> None:
    """Grava num arquivo temporário e troca no fim: nunca deixa uma saída pela metade."""
    extensao = destino.suffix.lower()
    if extensao not in EXTENSOES_DE_SAIDA:
        raise ErroDePlanilha(f"Formato de saída não suportado: {destino.name} (use .xlsx ou .csv)")
    destino.parent.mkdir(parents=True, exist_ok=True)
    descritor, temporario = tempfile.mkstemp(suffix=extensao, dir=destino.parent)
    os.close(descritor)
    try:
        (_gravar_csv if extensao == ".csv" else _gravar_xlsx)(pasta, Path(temporario))
        os.replace(temporario, destino)
    except BaseException:
        Path(temporario).unlink(missing_ok=True)
        raise
